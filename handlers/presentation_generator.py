"""
Docura.kz — Генератор интерактивных презентаций PowerPoint (.pptx)
Создает методические презентации для уроков и занятий по стандартам МОН РК.
"""
import os
import tempfile
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


class PresentationGenerator:
    def __init__(self, api_key: str = ""):
        self.api_key = api_key or os.getenv("ANTHROPIC_API_KEY", "")

    async def generate_lesson_presentation(
        self,
        topic: str,
        subject: str,
        grade: str,
        lang: str = "ru",
        teacher_name: str = "",
        school: str = ""
    ) -> str:
        """
        Генерирует контент через Claude и собирает файл .pptx.
        Возвращает путь к сгенерированному файлу .pptx.
        """
        slides_data = await self._generate_slides_content(topic, subject, grade, lang, teacher_name, school)
        return self._build_pptx(slides_data, lang)

    async def _generate_slides_content(
        self,
        topic: str,
        subject: str,
        grade: str,
        lang: str,
        teacher_name: str,
        school: str
    ) -> list:
        """Обращается к AI для составления методической структуры презентации."""
        import anthropic

        prompt_lang = "казахском" if lang == "kz" else "русском"
        prompt = f"""Ты — ведущий методист образования Казахстана.
Составь план презентации к уроку на {prompt_lang} языке строго по обновленным стандартам МОН РК.

Тема урока: {topic}
Предмет: {subject}
Класс / группа: {grade}
ФИО педагога: {teacher_name or 'Педагог'}
Школа / организация: {school or 'Общеобразовательная школа'}

Требуется ровно 6-7 слайдов:
1. Титульный слайд (тема, предмет, класс, учитель)
2. Цели обучения (Сабақтың мақсаттары / Цели урока по ГОСО РК)
3. Ой қозғау / Актуализация знаний (2-3 проблемных вопроса для класса)
4. Негізгі ұғымдар / Теоретический материал (3-4 ключевых тезиса с пояснениями)
5. Жұптық / топтық тапсырмалар / Практические задания (дифференцированные задачи)
6. Қалыптастырушы бағалау / Дескрипторы (критерии оценивания)
7. Рефлексия және үй тапсырмасы / Рефлексия и домашнее задание (метод светофора/лестницы успеха)

Ответь ТОЛЬКО валидным JSON списком объектов без markdown блоков. Формат:
[
  {{
    "title": "Заголовок слайда",
    "subtitle": "Подзаголовок или категория",
    "bullets": ["Текст пункта 1", "Текст пункта 2", "Текст пункта 3"],
    "footer": "Подпись внизу (например: МОН РК • Дескрипторы)"
  }}
]"""

        try:
            client = anthropic.AsyncAnthropic(api_key=self.api_key)
            resp = await client.messages.create(
                model="claude-haiku-4-5-20251001",
                max_tokens=1500,
                messages=[{"role": "user", "content": prompt}]
            )
            raw = resp.content[0].text.strip()
            import re
            raw = re.sub(r"^```[a-z]*", "", raw, flags=re.MULTILINE)
            raw = re.sub(r"```$", "", raw, flags=re.MULTILINE).strip()
            data = json.loads(raw)
            if isinstance(data, list) and len(data) >= 4:
                return data
        except Exception as e:
            print(f"Presentation AI gen error: {e}")

        # Методический fallback шаблон, если API недоступен
        if lang == "kz":
            return [
                {
                    "title": topic,
                    "subtitle": f"{subject} • {grade}-сынып",
                    "bullets": [f"Мұғалім: {teacher_name or 'Педагог'}", f"Ұйым: {school or 'Мектеп'}", "ҚР Оқу-ағарту министрлігі стандарттары бойынша"],
                    "footer": "Docura.kz • Мұғалімнің көмекшісі"
                },
                {
                    "title": "Сабақтың мақсаттары",
                    "subtitle": "Оқу бағдарламасына сәйкес",
                    "bullets": [
                        f"{topic} негізгі ұғымдарымен танысу және тәжірибеде қолдану",
                        "Логикалық ойлау және талдау дағдыларын дамыту",
                        "Өзара сыйластық пен ұжымдық жұмысқа баулу"
                    ],
                    "footer": "ҚР ГОСО • 2026"
                },
                {
                    "title": "Ой қозғау",
                    "subtitle": "Өткенді қайталау және жаңа сабаққа көпір",
                    "bullets": [
                        "Бұл тақырып біздің күнделікті өмірімізде қайда кездеседі?",
                        "Өткен тақырыппен қандай байланысы бар?",
                        "Бүгінгі сабақта не үйренгіміз келеді?"
                    ],
                    "footer": "Интерактивті талқылау"
                },
                {
                    "title": "Жаңа сабақты меңгеру",
                    "subtitle": "Негізгі теориялық түсініктер",
                    "bullets": [
                        "1-ереже: Негізгі анықтамалар мен ұғымдар жүйесі",
                        "2-ереже: Формулалар мен тәжірибелік мысалдар",
                        "3-ереже: Қателіктерді талдау және дұрыс қолдану алгоритмі"
                    ],
                    "footer": "Теория және практика"
                },
                {
                    "title": "Тәжірибелік тапсырмалар",
                    "subtitle": "Деңгейлік тапсырмалар (A, B, C)",
                    "bullets": [
                        "А деңгейі: Негізгі ұғымды бекітуге арналған жаттығу",
                        "В деңгейі: Алған білімді жаңа жағдайда қолдану",
                        "С деңгейі: Шығармашылық және ізденіс тапсырмасы"
                    ],
                    "footer": "Дифференциация"
                },
                {
                    "title": "Рефлексия және үй тапсырмасы",
                    "subtitle": "Сабақты қорытындылау",
                    "bullets": [
                        "«Табыс баспалдағы»: бүгін мен не білдім, не қиын болды?",
                        f"Үй тапсырмасы: {topic} бойынша оқулық параграфы мен №1-3 есептер",
                        "Кері байланыс: сұрақтар мен ұсыныстар"
                    ],
                    "footer": "Docura.kz • Сабақ аяқталды"
                }
            ]
        else:
            return [
                {
                    "title": topic,
                    "subtitle": f"{subject} • {grade} класс",
                    "bullets": [f"Учитель: {teacher_name or 'Педагог'}", f"Организация: {school or 'Школа'}", "По стандартам Министерства просвещения РК"],
                    "footer": "Docura.kz • Помощник педагога"
                },
                {
                    "title": "Цели и задачи урока",
                    "subtitle": "В соответствии с учебной программой",
                    "bullets": [
                        f"Усвоить ключевые понятия темы: {topic}",
                        "Сформировать умение применять знания на практике",
                        "Развивать критическое мышление и функциональную грамотность"
                    ],
                    "footer": "ГОСО РК • КСП"
                },
                {
                    "title": "Актуализация знаний",
                    "subtitle": "Мозговой штурм и связь с предыдущей темой",
                    "bullets": [
                        "Где в реальной жизни мы сталкиваемся с этой темой?",
                        "Какие базовые понятия нам понадобятся сегодня?",
                        "Что станет главным результатом нашего урока?"
                    ],
                    "footer": "Интерактивная разминка"
                },
                {
                    "title": "Изучение нового материала",
                    "subtitle": "Ключевые понятия и правила",
                    "bullets": [
                        "Тезис 1: Базовое определение и структура понятия",
                        "Тезис 2: Практический пример и алгоритм решения",
                        "Тезис 3: Частые ошибки и как их избежать"
                    ],
                    "footer": "Теория и формулы"
                },
                {
                    "title": "Практическая работа",
                    "subtitle": "Дифференцированные задания (Уровни A, B, C)",
                    "bullets": [
                        "Уровень А: Базовое применение алгоритма (для всех учащихся)",
                        "Уровень B: Задача с несколькими действиями и анализом",
                        "Уровень C: Нестандартная творческая задача повышенной сложности"
                    ],
                    "footer": "Формативное оценивание"
                },
                {
                    "title": "Рефлексия и домашнее задание",
                    "subtitle": "Подведение итогов",
                    "bullets": [
                        "Приём «Светофор»: зелёный (всё понятно), жёлтый (есть вопросы), красный (нужна помощь)",
                        f"Домашнее задание: изучить материал по теме «{topic}», выполнить упражнения",
                        "Спасибо за активную работу на уроке!"
                    ],
                    "footer": "Docura.kz • Урок завершён"
                }
            ]

    def _build_pptx(self, slides_data: list, lang: str) -> str:
        """Формирует файл .pptx с элегантным оформлением и слайдами 16:9."""
        prs = Presentation()
        # Стандарт 16:9 widescreen
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # Цвета
        NAVY = RGBColor(15, 31, 61)
        ROYAL_BLUE = RGBColor(37, 99, 235)
        DARK_TEXT = RGBColor(30, 41, 59)
        LIGHT_BG = RGBColor(248, 250, 252)
        WHITE = RGBColor(255, 255, 255)
        MUTED = RGBColor(100, 116, 139)

        blank_slide_layout = prs.slide_layouts[6]

        for i, s_info in enumerate(slides_data):
            slide = prs.slides.add_slide(blank_slide_layout)

            # Фон слайда
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg.line.fill.background()

            is_title = (i == 0)
            if is_title:
                bg.fill.solid()
                bg.fill.fore_color.rgb = NAVY

                # Декоративная полоса
                accent_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.2), Inches(3.8))
                accent_strip.fill.solid()
                accent_strip.fill.fore_color.rgb = ROYAL_BLUE
                accent_strip.line.fill.background()

                # Заголовок
                tb = slide.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(10.5), Inches(2.2))
                tf = tb.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = s_info.get("title", "Тема урока")
                p.font.size = Pt(40)
                p.font.bold = True
                p.font.color.rgb = WHITE

                # Подзаголовок
                p2 = tf.add_paragraph()
                p2.text = s_info.get("subtitle", "")
                p2.font.size = Pt(22)
                p2.font.color.rgb = RGBColor(96, 165, 250)
                p2.space_before = Pt(14)

                # Пункты / информация об учителе
                bullets_tb = slide.shapes.add_textbox(Inches(1.8), Inches(4.2), Inches(10.5), Inches(2.2))
                btf = bullets_tb.text_frame
                btf.word_wrap = True
                for b_idx, bullet in enumerate(s_info.get("bullets", [])):
                    bp = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
                    bp.text = f"•  {bullet}"
                    bp.font.size = Pt(16)
                    bp.font.color.rgb = RGBColor(226, 232, 240)
                    bp.space_before = Pt(6)

            else:
                # Обычный контентный слайд
                bg.fill.solid()
                bg.fill.fore_color.rgb = LIGHT_BG

                # Верхняя плашка шапки
                top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.3))
                top_bar.fill.solid()
                top_bar.fill.fore_color.rgb = NAVY
                top_bar.line.fill.background()

                # Заголовок слайда
                title_tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(10.0), Inches(0.9))
                ttf = title_tb.text_frame
                ttf.word_wrap = True
                tp = ttf.paragraphs[0]
                tp.text = s_info.get("title", "")
                tp.font.size = Pt(26)
                tp.font.bold = True
                tp.font.color.rgb = WHITE

                if s_info.get("subtitle"):
                    stp = ttf.add_paragraph()
                    stp.text = s_info.get("subtitle", "")
                    stp.font.size = Pt(13)
                    stp.font.color.rgb = RGBColor(147, 197, 253)

                # Номер слайда в правом углу
                num_tb = slide.shapes.add_textbox(Inches(11.2), Inches(0.3), Inches(1.5), Inches(0.6))
                np = num_tb.text_frame.paragraphs[0]
                np.alignment = PP_ALIGN.RIGHT
                np.text = f"{i + 1} / {len(slides_data)}"
                np.font.size = Pt(16)
                np.font.bold = True
                np.font.color.rgb = RGBColor(147, 197, 253)

                # Карточка с тезисами
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
                card.fill.solid()
                card.fill.fore_color.rgb = WHITE
                card.line.color.rgb = RGBColor(226, 232, 240)
                card.line.width = Pt(1.5)

                # Контент карточки
                content_tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(4.6))
                ctf = content_tb.text_frame
                ctf.word_wrap = True

                bullets = s_info.get("bullets", [])
                for b_idx, bullet in enumerate(bullets):
                    bp = ctf.paragraphs[0] if b_idx == 0 else ctf.add_paragraph()
                    bp.text = f"✔  {bullet}"
                    bp.font.size = Pt(18)
                    bp.font.color.rgb = DARK_TEXT
                    bp.space_before = Pt(16 if len(bullets) <= 4 else 10)
                    bp.space_after = Pt(4)

                # Футер
                footer_text = s_info.get("footer", "Docura.kz • ҚР стандарты")
                foot_tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.5))
                fp = foot_tb.text_frame.paragraphs[0]
                fp.text = footer_text
                fp.font.size = Pt(11)
                fp.font.color.rgb = MUTED

        # Сохранение во временный файл
        fd, out_path = tempfile.mkstemp(suffix=".pptx", prefix="Docura_Presentation_")
        os.close(fd)
        prs.save(out_path)
        return out_path
