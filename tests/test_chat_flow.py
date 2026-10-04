"""
Сквозной тест ИИ-агента в чате без интернета: поддельные сервера Anthropic и OpenAI-API
запускаются в этом же процессе. Запуск:  python tests/test_chat_flow.py
Проверяет: время и приветствия, память, «печатает», сбор недостающих данных в разговоре,
выбор языка, локализацию шапок таблиц в Word и лимит бесплатных документов.
"""
import asyncio, json, os, sys, tempfile, threading, types, io, re
from http.server import BaseHTTPRequestHandler, HTTPServer
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

CALLS = {"anthropic": [], "openai": []}
CHAT_REPLY = {"reply": "Добрый день, Айгуль! Чем помочь?", "action": "none", "doc_type": None,
              "remember": ["преподаёт математику в 7А"], "forget": False}

RU_DOC = """КРАТКОСРОЧНЫЙ ПЛАН УРОКА
Предмет: Математика
Класс: 7А
ЦЕЛИ ОБУЧЕНИЯ:
- Қосу ережесін қолданады
| Этап урока | Деятельность педагога и ученика | Оценивание и ресурсы |
| Начало (5 мин) | Мұғалім сұрақ қояды | Ауызша кері байланыс |
РЕФЛЕКСИЯ:
Светофор"""


class Fake(BaseHTTPRequestHandler):
    def log_message(self, *a): pass

    def _send(self, obj):
        b = json.dumps(obj).encode()
        self.send_response(200); self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(b))); self.end_headers(); self.wfile.write(b)

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if self.path.endswith("/chat/completions"):
            CALLS["openai"].append(body)
            return self._send({"id": "x", "object": "chat.completion", "created": 0, "model": "m",
                               "choices": [{"index": 0, "finish_reason": "stop",
                                            "message": {"role": "assistant", "content": json.dumps(CHAT_REPLY, ensure_ascii=False)}}],
                               "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2}})
        prompt = body["messages"][0]["content"]
        prompt = prompt if isinstance(prompt, str) else prompt[0].get("text", "")
        CALLS["anthropic"].append(prompt[:60])
        if "Определи, является ли" in prompt:
            text = json.dumps({"doc_type": "lesson_plan"})
        elif "Ты помогаешь собрать данные" in prompt:
            msg = re.search(r'Сообщение пользователя:\n"(.*)"\n', prompt, re.S).group(1)
            ans, lang, intent = {}, None, "continue"
            if "7А" in msg: ans["subject_class"] = "Математика, 7А"
            m = re.search(r"тема[: ]+([^,.]+)", msg, re.I)
            if m: ans["topic"] = m.group(1).strip()
            if "45" in msg: ans["date"] = "завтра, 45 минут"
            if "казах" in msg: lang = "kz"
            if "отмен" in msg.lower(): intent = "cancel"
            if "погода" in msg: intent = "other"
            text = json.dumps({"intent": intent, "answers": ans, "doc_lang": lang, "student_hint": None}, ensure_ascii=False)
        elif "Оцени этот официальный документ" in prompt:
            text = "92"
        else:
            text = RU_DOC
        self._send({"id": "m", "type": "message", "role": "assistant", "model": body["model"],
                    "content": [{"type": "text", "text": text}], "stop_reason": "end_turn",
                    "stop_sequence": None, "usage": {"input_tokens": 1, "output_tokens": 1}})


server = HTTPServer(("127.0.0.1", 0), Fake)
threading.Thread(target=server.serve_forever, daemon=True).start()
port = server.server_address[1]
os.environ.update({"ANTHROPIC_BASE_URL": f"http://127.0.0.1:{port}", "CONCIERGE_BASE_URL": f"http://127.0.0.1:{port}",
                   "CONCIERGE_API_KEY": "k", "ANTHROPIC_API_KEY": "k", "TELEGRAM_TOKEN": "1:x",
                   "DB_PATH": tempfile.mktemp(suffix=".db")})

from database import Database
from handlers import concierge as C
from handlers.concierge import ConciergeHandler


class Msg:
    def __init__(self, chat_id, text=""):
        self.chat_id, self.text, self.sent, self.docs = chat_id, text, [], []
    async def reply_text(self, text, **kw):
        self.sent.append((text, kw)); return self
    async def reply_document(self, document, filename=None, **kw):
        self.docs.append((filename, document.read(), kw))
    async def delete(self): pass
    async def edit_text(self, *a, **k): pass


class Bot:
    username = "docura_test_bot"
    def __init__(self): self.actions = []
    async def send_chat_action(self, chat_id, action): self.actions.append(action)
    async def send_message(self, *a, **k): pass


def upd(uid, text):
    m = Msg(uid, text)
    return types.SimpleNamespace(message=m, effective_user=types.SimpleNamespace(id=uid), callback_query=None), m


def fake_now(hour):
    real = C.now_local()
    return lambda: real.replace(hour=hour, minute=30)


PASS, FAIL = [], []
def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("✅ " if cond else "❌ ") + name + (f"  {extra}" if extra and not cond else ""))


async def main():
    db = Database(); await db.init()
    await db.upsert_user(1, {"name": "Айгуль Н.", "lang": "ru", "role": "teacher", "school": "Лицей 5",
                            "subject": "Математика", "classes": "7А", "subscribed": 1, "director": "Иванов И.И."})
    h = ConciergeHandler(db, "k")
    ctx = types.SimpleNamespace(bot=Bot(), user_data={}, application=types.SimpleNamespace(bot_data={}))

    # ── 1. Просьба сделать документ: агент пишет, чего не хватает ──
    u, m = upd(1, "сделай КСП по математике 7А")
    await h.handle_text(u, ctx)
    txt = m.sent[-1][0]
    check("агент просит недостающие данные (тема/длительность/язык)",
          "не хватает" in txt and "Тема урока" in txt and "длительность" in txt.lower() and "язык" in txt.lower(), txt)
    check("предмет и класс (названы в сообщении) повторно НЕ спрашиваются", "Предмет и класс" not in txt)
    check("необязательные цели обучения не спрашиваются", "Цели обучения" not in txt)
    check("кнопки выбора языка показаны", any("cg_lang_kz" in str(m.sent[-1][1].get("reply_markup")) for _ in [0]))
    check("индикатор «печатает» отправлялся", "typing" in ctx.bot.actions, str(ctx.bot.actions))

    # ── 2. Пользователь отвечает свободным текстом -> генерация на казахском ──
    u, m = upd(1, "Тема: Сложение дробей, 45 минут, на казахском")
    await h.handle_text(u, ctx)
    check("документ отправлен после ответа", len(m.docs) == 1, str([s[0] for s in m.sent]))
    if m.docs:
        from docx import Document
        doc = Document(io.BytesIO(m.docs[0][1]))
        flat = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        check("шапка таблицы переведена на казахский", "Сабақ кезеңі" in flat and "Этап урока" not in flat, flat[:300])
        check("заголовки разделов переведены на казахский", "ОҚУ МАҚСАТТАРЫ" in flat and "ЦЕЛИ ОБУЧЕНИЯ" not in flat)
        check("подписи полей шапки переведены", "Пән" in flat and "Предмет:" not in flat)
        check("подпись в конце на казахском", "Басшы" in flat)
        check("заголовок документа на казахском", "ҚМЖ" in flat.upper() and "КРАТКОСРОЧНЫЙ" not in flat.upper(), flat[:120])
    st = await h._load_state(1)
    check("задача очищена после генерации", st["doc_task"] is None)

    # ── 3. Приветствие: после генерации разговор «идёт» -> не здороваемся ──
    CHAT_REPLY["reply"] = "Добрый день, Айгуль! Рада помочь, что дальше?"
    u, m = upd(1, "спасибо, а подскажи как лучше провести рефлексию?")
    await h.handle_text(u, ctx)
    reply = m.sent[-1][0]
    check("в идущем разговоре приветствие убрано", not reply.lower().startswith(("добрый", "здравствуйте", "привет")), reply)
    sysmsg = CALLS["openai"][-1]["messages"][0]["content"]
    check("в промпт передано местное время и запрет здороваться", "ТЕКУЩЕЕ ВРЕМЯ" in sysmsg and "НЕ здоровайся" in sysmsg)
    check("история передаётся модели (память)", len(CALLS["openai"][-1]["messages"]) > 3)
    st = await h._load_state(1)
    check("факт сохранён в долгой памяти", any("математику" in f for f in st["facts"]), str(st["facts"]))

    # ── 4. Новый разговор ночью: правильное приветствие по местному времени ──
    await db.update_agent_context(2, {})
    await db.upsert_user(2, {"name": "Ермек", "lang": "ru", "role": "teacher", "school": "Школа"})
    for hour, expect in ((3, "Здравствуйте"), (9, "Доброе утро"), (14, "Добрый день"), (20, "Добрый вечер")):
        C.now_local = fake_now(hour)
        await db.update_agent_context(2, {"concierge": {}})
        CHAT_REPLY["reply"] = "Добрый день, Ермек! Чем помочь?"
        u, m = upd(2, "привет")
        await h.handle_text(u, ctx)
        r = m.sent[-1][0]
        check(f"{hour}:30 -> «{expect}»", r.startswith(expect), r)
    import importlib; importlib.reload(C); h = ConciergeHandler(db, "k")

    # ── 5. Длинная пауза -> здороваемся снова ──
    st = await h._load_state(1)
    for t in st["history"]: t["ts"] = "2020-01-01T10:00:00+05:00"
    await h._save_state(1, st)
    CHAT_REPLY["reply"] = "Добрый вечер, Айгуль! Как дела?"
    u, m = upd(1, "я снова тут")
    await h.handle_text(u, ctx)
    check("после паузы бот здоровается заново", m.sent[-1][0].split(",")[0] in ("Здравствуйте", "Доброе утро", "Добрый день", "Добрый вечер"), m.sent[-1][0])

    # ── 6. Отмена и посторонние сообщения во время сбора ──
    u, m = upd(1, "сделай КСП по математике"); await h.handle_text(u, ctx)
    u, m = upd(1, "отмена, не надо"); await h.handle_text(u, ctx)
    st = await h._load_state(1)
    check("отмена сбрасывает задачу", st["doc_task"] is None)
    u, m = upd(1, "сделай КСП"); await h.handle_text(u, ctx)
    CHAT_REPLY["reply"] = "Сегодня солнечно. Кстати, по КСП мне ещё нужны данные."
    u, m = upd(1, "какая сегодня погода?"); await h.handle_text(u, ctx)
    st = await h._load_state(1)
    check("посторонний вопрос не ломает задачу (она сохранена)", st["doc_task"] is not None)
    check("модель знает о незавершённом документе", "НЕЗАВЕРШЕНА" in CALLS["openai"][-1]["messages"][0]["content"])

    # ── 7. Лимит бесплатных документов ──
    await db.upsert_user(3, {"name": "Бесплатный", "lang": "ru", "role": "teacher", "school": "Ш", "subscribed": 0, "free_used": 999})
    u, m = upd(3, "сделай КСП по математике 7А"); await h.handle_text(u, ctx)
    u, m = upd(3, "Тема: Дроби, 45 минут, на казахском"); await h.handle_text(u, ctx)
    check("при исчерпанном лимите документ не создаётся, показывается тариф",
          len(m.docs) == 0 and any("PRO" in s[0] or "лимит" in s[0].lower() or "Бесплатные" in s[0] for s in m.sent), str([s[0][:60] for s in m.sent]))

    print(f"\n{len(PASS)} прошло, {len(FAIL)} упало")
    return 1 if FAIL else 0

if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
