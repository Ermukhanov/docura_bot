"""Проверка полноты схем документов на ru/kz/en и работы Word-генератора. Запуск: python tests/test_schemas.py"""
import os, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault("DB_PATH", tempfile.mktemp(suffix=".db"))

from handlers import documents as D, doc_schemas as S, doc_questions as Q
from handlers.word_generator import generate_word
from docx import Document

fails = []
def check(name, cond, extra=""):
    print(("✅ " if cond else "❌ ") + name + (f"  {extra}" if extra and not cond else ""))
    if not cond: fails.append(name)

get = lambda dt, lg: Q.get_questions(dt, lg, D.DOC_QUESTIONS, D.REGISTRY_QUESTIONS)
probs = S.coverage_problems(D.DOC_NAMES, get)
check("для каждого документа есть название, вопросы и подписи на ru/kz/en", not probs, "; ".join(probs[:8]))
check("у всех ключей вопросов есть английский текст", not Q.all_keys_translated(D.DOC_QUESTIONS, D.REGISTRY_QUESTIONS),
      str(Q.all_keys_translated(D.DOC_QUESTIONS, D.REGISTRY_QUESTIONS)))

real = set()
for cats in (D.CAT_DOCS, D.CAT_DOCS_KG, D.CAT_DOCS_COMMON):
    for v in cats.values(): real.update(v)
real.update(["kindergarten_cycle_schedule", "development_monitoring"])
check("схема есть у каждого типа документа", real <= set(S.SCHEMAS), str(real - set(S.SCHEMAS)))

# одинаковые наборы полей на всех языках
bad = []
for dt in real:
    keys = {lg: [q["key"] for q in get(dt, lg)] for lg in S.LANGS}
    if keys["ru"] and (keys["kz"] != keys["ru"] or keys["en"] != keys["ru"]): bad.append(dt)
check("набор вопросов одинаков на всех языках", not bad, str(bad))

# инструкция по структуре и локализация для КАЖДОГО документа на КАЖДОМ языке
for lg in ("kz", "en"):
    leaks = []
    for dt in sorted(S.SCHEMAS):
        sch = S.SCHEMAS[dt]
        lines = [S.label(k, "ru").upper() + ":" for k in sch["headings"]]
        for cols in sch["tables"]:
            lines.append("| " + " | ".join(S.label(c, "ru") for c in cols) + " |")
            lines.append("| какой-то текст | ещё текст |")
        lines += [f"{S.label(k, 'ru')}: значение" for k in sch["fields"]]
        out = S.localize_labels("\n".join(lines), lg)
        left = S.find_foreign_labels(out, lg)
        if left: leaks.append((dt, left[:3]))
    check(f"после локализации ни одной подписи структуры на другом языке ({lg})", not leaks, str(leaks[:4]))

check("английские формулировки документов без кириллицы",
      not [dt for dt, p in S.PHRASES.items() if any(S.has_cyrillic(s) for s in p["en"])])
check("инструкция по структуре строится для каждого документа и языка",
      all("ЯЗЫК СТРУКТУРЫ" in S.build_structure_instruction(dt, lg) for dt in S.SCHEMAS for lg in S.LANGS))

# Word-генератор: циклограмма / мониторинг / карта развития на всех языках
RU_ONLY = ["Организация", "Группа", "Период", "Воспитатель", "Планируемое содержание недели",
           "Понедельник", "Режимный момент", "Учесть при планировании"]
for lg in ("kz", "en"):
    data = {"organization": "Сад", "group": "Ясли", "period": "01-05.09", "week_topic": "Транспорт",
            "educator_name": "Иванова", "events": "утренник", "lang": lg}
    f = generate_word("", "Циклограмма", "Иванова", "Петрова", cycle_data=data, lang=lg)
    d = Document(f); txt = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells) + "\n".join(p.text for p in d.paragraphs)
    check(f"циклограмма ({lg}) без русских подписей структуры", not any(w in txt for w in RU_ONLY), str([w for w in RU_ONLY if w in txt]))
    f = generate_word("", "Мониторинг", "И", "П", monitoring_data={**data, "children": ["А", "Б"], "rows": []}, lang=lg)
    d = Document(f); txt = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    check(f"мониторинг ({lg}) без русских подписей", "ФИО ребенка" not in txt and "Организация" not in txt and "Физическое развитие" not in txt)
    if lg == "en":
        check("мониторинг (en) не падает в казахский", "Баланың" not in txt and "Child's full name" in txt)
    f = generate_word("", "Карта", "И", "П", monitoring_data={**data, "child_name": "Али"}, registry_doc_type="kg_individual_development_card", lang=lg)
    check(f"карта развития ({lg}) формируется", os.path.exists(f))

print(f"\n{'ВСЁ ОК' if not fails else 'ЕСТЬ ОШИБКИ: ' + str(len(fails))}")
sys.exit(1 if fails else 0)
